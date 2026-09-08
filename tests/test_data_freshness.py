import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from sirena.data_loader import (
    MONTHLY_COLUMNS, DataFreshnessError, DataLoader, load_model_data,
    require_observations, forecast_source_manifest, validate_forecast_cache,
)


@pytest.fixture
def sources(tmp_path):
    dates = pd.date_range('2025-01-01', periods=19, freq='MS')
    raw = pd.DataFrame({'Date': dates.strftime('%d.%m.%Y')})
    for n, key in enumerate(MONTHLY_COLUMNS):
        raw[key] = 100 + n + np.arange(len(dates)) / 100
    for col in ('usd_nom_i', 'Ki', 'Ruonia', 'Ki_i'):
        raw[col] = 10.0
    raw.to_csv(tmp_path / 'inflation_data.csv', sep=';', decimal=',', index=False)
    sa = pd.DataFrame([
        {'Код': n, 'Товар': name, 'Дата': date.strftime('%d.%m.%Y'), 'Значение': 99+n/10}
        for date in dates for n, name in enumerate(MONTHLY_COLUMNS.values(), 1)
    ])
    sa.to_csv(tmp_path / 'sa_fl.csv', sep=';', decimal=',', index=False)
    return tmp_path


def test_raw_uses_canonical_even_if_alternative_is_stale(sources):
    (sources / 'infl_kbr.csv').write_text('invalid stale alternative')
    frame = DataLoader(sources).load_monthly_kbr()
    assert frame.index.max() == pd.Timestamp('2026-07-01')
    assert frame.iloc[-1, 0] == pytest.approx(100.18)
    assert frame.attrs['input_contract']['representation'] == 'raw'


def test_sa_remains_homogeneous_and_cutoff_excludes_future(sources):
    frame = load_model_data('sa', sources, cutoff='2026-03-15')
    assert frame.index.max() == pd.Timestamp('2026-03-01')
    assert (frame.iloc[:, 0] == 99.1).all()
    assert frame.attrs['input_contract']['representation'] == 'sa'


def test_sa_lag_is_error_never_raw_splice(sources):
    path = sources / 'sa_fl.csv'
    frame = pd.read_csv(path, sep=';', decimal=',')
    frame = frame[frame['Дата'] != '01.07.2026']
    frame.to_csv(path, sep=';', decimal=',', index=False)
    with pytest.raises(DataFreshnessError, match='2026-07'):
        load_model_data('sa', sources)
    assert load_model_data('raw', sources).index.max().month == 7


@pytest.mark.parametrize('missing_month', ['01.07.2026', '01.05.2026'])
def test_one_missing_component_or_lag_blocks(sources, missing_month):
    path = sources / 'inflation_data.csv'
    raw = pd.read_csv(path, sep=';', decimal=',')
    raw.loc[raw.Date == missing_month, 'Serv'] = np.nan
    raw.to_csv(path, sep=';', decimal=',', index=False)
    with pytest.raises(DataFreshnessError):
        load_model_data('raw', sources)


def test_changed_source_invalidates_cache_even_same_latest_date(sources):
    payload = {'source_manifest': forecast_source_manifest(sources)}
    validate_forecast_cache(payload, sources)
    path = sources / 'inflation_data.csv'
    path.write_bytes(path.read_bytes().replace(b'100,01', b'100,02'))
    with pytest.raises(DataFreshnessError, match='inflation_data'):
        validate_forecast_cache(payload, sources)


def test_legacy_cache_rejected(sources):
    with pytest.raises(DataFreshnessError):
        validate_forecast_cache({}, sources)


def test_partial_future_month_cannot_silently_move_origin_back(sources):
    path = sources / 'inflation_data.csv'
    raw = pd.read_csv(path, sep=';', decimal=',')
    raw.loc[len(raw)] = {'Date': '01.08.2026', 'mom': 100.5}
    raw.to_csv(path, sep=';', decimal=',', index=False)
    with pytest.raises(DataFreshnessError, match='2026-08'):
        load_model_data('raw', sources)


def test_missing_exact_external_forecast_origin_is_not_relabelled():
    from sirena.models.micro_statsmodels_external import MicroStatsmodelsExternalForecaster
    model = MicroStatsmodelsExternalForecaster()
    model._is_fitted = True
    model._last_train_date = pd.Timestamp('2026-07-01')
    model._forecasts = pd.DataFrame({pd.Timestamp('2026-05-01'): [100.5]})
    with pytest.raises(DataFreshnessError, match='required cutoff 2026-07'):
        model.forecast(1)


def test_internal_micro_is_truncated_and_stale_rejected(monkeypatch):
    from sirena.models.microcomponent import MicrocomponentForecaster
    model = MicrocomponentForecaster()
    micro = pd.DataFrame({101: [0.1, 0.2]}, index=pd.to_datetime(['2026-06-01', '2026-08-01']))
    monkeypatch.setattr(model, '_load_data', lambda _: micro)
    with pytest.raises(DataFreshnessError, match='2026-07'):
        model.fit(pd.DataFrame({'mom': [100.5]}, index=pd.to_datetime(['2026-07-01'])))


def test_production_origin_requires_correct_observed_lag_window(tmp_path):
    from sirena.macro_features import validate_production_origin
    dates = pd.date_range('2026-01-01', periods=7, freq='MS')
    proxy = pd.DataFrame({'Date': dates.strftime('%d.%m.%Y'), 'Torg': 101., 'pp': 102.})
    proxy.loc[5:, 'pp'] = np.nan  # lag3 does not need June/July for August h1.
    proxy.to_csv(tmp_path / 'infostat.csv', index=False, sep=';', decimal=',')
    validate_production_origin('2026-07-01', tmp_path)
    proxy.loc[6, 'Torg'] = np.nan
    proxy.to_csv(tmp_path / 'infostat.csv', index=False, sep=';', decimal=',')
    with pytest.raises(DataFreshnessError):
        validate_production_origin('2026-07-01', tmp_path)


def test_production_and_backtest_share_latest_raw_values():
    from scripts.precompute_forecasts import _load_forecast_input_data
    from scripts.backtest_framework import BacktestRunner
    live = _load_forecast_input_data()
    runner = BacktestRunner.__new__(BacktestRunner)
    runner._prepare_data()
    pd.testing.assert_frame_equal(live, runner.df_ridge)
    runner.horizon = 1
    train, _, cutoff = runner._train_test_split(live.index.max())
    assert train.index.max() == cutoff
    assert (train.index < live.index.max()).all()


def test_forecast_storage_rejects_partial_nonfinite_output():
    from scripts.precompute_forecasts import _store_forecast
    with pytest.raises(DataFreshnessError):
        _store_forecast({'forecasts': {}}, 'Example', [0.2, np.nan])


def test_raw_model_rejects_tagged_sa(sources):
    from sirena.models.ridge import RidgeForecaster
    with pytest.raises(ValueError, match='expected raw, received sa'):
        RidgeForecaster().fit(load_model_data('sa', sources))


def test_policy_json_is_not_mistaken_for_forecast_cache(tmp_path):
    from scripts.generate_forecast_table import load_forecasts
    path = tmp_path / 'send_ready_policy_trajectory.json'
    path.write_text(json.dumps({'mom_pp': [0.5]}))
    assert load_forecasts(path)['mom_pp'] == [0.5]


def test_sa_importer_cross_checks_db_and_region(tmp_path):
    import openpyxl
    from scripts.import_sa_vintage import extract_vintage
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = 'mom_sa'
    ws.cell(1, 7, 'Кабардино-Балкарская Республика')
    ws.cell(1, 11, 7)
    ws.cell(3, 2, 'Темп роста м/м, с устранением сезонности (momsa)')
    date = pd.Timestamp('2026-07-01').to_pydatetime()
    ws.cell(4, 4, date)
    db = workbook.create_sheet('db')
    db.append(['Date', 'item', 'region', 'raw', 'mom_sa'])
    for code in (1, 2, 3, 4):
        ws.cell(4+code, 2, code)
        ws.cell(4+code, 3, f'item{code}')
        ws.cell(4+code, 4, 101.)
        db.append([date, code, 7, 100., 101.])
    path = tmp_path / 'sa.xlsx'
    workbook.save(path)
    assert extract_vintage(path)[2] == 4
    db.cell(2, 5, 102.)
    workbook.save(path)
    with pytest.raises(ValueError, match='mismatch'):
        extract_vintage(path)
    ws.cell(1, 11, 8)
    workbook.save(path)
    with pytest.raises(ValueError, match='region 7'):
        extract_vintage(path)
