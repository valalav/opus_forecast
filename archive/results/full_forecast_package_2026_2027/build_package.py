#!/usr/bin/env python3
"""Build forecast CSV, DOCX narrative, and ZIP package."""
from __future__ import annotations

import csv
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "archive" / "results" / "full_forecast_package_2026_2027"
CSV_PATH = OUT / "forecast_2026_2027_mom_yoy.csv"
DOCX_PATH = OUT / "forecast_2026_2027_explanation.docx"
ZIP_PATH = OUT / "sirena_kbr_forecast_package_2026_2027.zip"

FORECAST_ROWS = [
    ("01.01.2026", "101,57", "106,04"),
    ("01.02.2026", "101,10", "105,79"),
    ("01.03.2026", "100,54", "105,50"),
    ("01.04.2026", "99,72", "104,65"),
    ("01.05.2026", "99,90", "104,62"),
    ("01.06.2026", "100,00", "104,66"),
    ("01.07.2026", "99,90", "104,18"),
    ("01.08.2026", "99,85", "104,62"),
    ("01.09.2026", "100,45", "104,06"),
    ("01.10.2026", "101,15", "104,70"),
    ("01.11.2026", "100,25", "104,80"),
    ("01.12.2026", "100,30", "104,81"),
    ("01.01.2027", "100,95", "104,17"),
    ("01.02.2027", "100,65", "103,71"),
    ("01.03.2027", "100,50", "103,67"),
    ("01.04.2027", "100,45", "104,43"),
    ("01.05.2027", "100,20", "104,74"),
    ("01.06.2027", "100,10", "104,84"),
    ("01.07.2027", "100,05", "105,00"),
    ("01.08.2027", "99,95", "105,11"),
    ("01.09.2027", "100,45", "105,11"),
    ("01.10.2027", "100,45", "104,38"),
    ("01.11.2027", "100,20", "104,33"),
    ("01.12.2027", "100,00", "104,02"),
]


def write_csv() -> None:
    with CSV_PATH.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(["Date", "MoM", "YoY"])
        writer.writerows(FORECAST_ROWS)


def add_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Date", "MoM", "YoY"]):
        table.rows[0].cells[i].text = h
    for row in FORECAST_ROWS:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v


def write_docx() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Прогноз ИПЦ КБР: 2026–2027")
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph(
        "Документ фиксирует обновлённый прогноз ИПЦ КБР после выхода факта за апрель 2026 "
        "и обновления недельных данных за три недели мая. Прогноз сочетает фактические значения "
        "января–апреля 2026, nowcast мая, оценку июньской продовольственной динамики, возможный "
        "возврат части снизившихся продовольственных цен, крепкий рубль на 21 мая и сценарное "
        "ослабление валютного фактора в дальнейшем."
    )

    doc.add_heading("Ключевые предпосылки", level=1)
    for item in [
        "Факт апреля 2026 составил 99,72, что заметно ниже прежнего прогноза 100,45.",
        "Майский прогноз установлен на 99,90: достигнутые цены на 18 мая относительно последней недели апреля дают около 99,93, но остаётся риск слабой последней недели.",
        "Июнь установлен на 100,00: продовольствие исторически слабо в июне, но значительная часть плодоовощного снижения уже реализована в апреле–мае.",
        "С июля месячные значения не менялись кардинально: сохранён плавный профиль с учётом летней сезонности, возможного возврата части продовольственных цен и будущего ослабления рубля.",
        "Июльская индексация услуг в 2026 году не закладывается; тарифно-услуговый эффект смещён на октябрь, но в сглаженной форме.",
        "Конец 2027 года настроен близко к целевым 4% YoY: декабрь 2027 = 104,02.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Прогноз", level=1)
    add_table(doc)

    doc.add_heading("Интерпретация", level=1)
    doc.add_paragraph(
        "Основное изменение профиля связано с апрельско-майской продовольственной дефляцией. "
        "Снижение огурцов, помидоров и яиц оказалось сильным, но по данным на 18 мая часть падения "
        "компенсируется ростом капусты, лука, круп, мясных изделий, отдельных непродовольственных "
        "товаров и услуг. Поэтому наиболее вероятный майский индекс сейчас ближе к 99,9, а не к "
        "более глубоким дефляционным сценариям."
    )
    doc.add_paragraph(
        "В июне сохраняется риск продовольственной слабости, однако большая часть раннего "
        "плодоовощного снижения уже реализована. Центральная оценка июня — около 100,0. Далее "
        "прогноз предполагает постепенную нормализацию и частичный возврат продовольственных цен, "
        "особенно после летнего минимума."
    )
    doc.add_paragraph(
        "Курс рубля на 21 мая около 71 за доллар является дезинфляционным фактором через импортные "
        "и торгуемые товары. Одновременно это может быть связано с падением импорта; в дальнейшем "
        "возможен частичный рост доллара, но неопределённость высокая, поэтому валютный разворот "
        "заложен мягко, без резкого переноса в месячные индексы."
    )

    doc.add_heading("Пояснения терминов для руководства", level=1)
    glossary = [
        ("Индекс MoM", "Индекс к предыдущему месяцу. Значение 100,20 означает рост цен на 0,20% за месяц; 99,90 — снижение на 0,10%."),
        ("YoY", "Индекс к тому же месяцу прошлого года. Значение 104,02 означает, что цены примерно на 4,02% выше, чем годом ранее."),
        ("п.п.", "Процентный пункт. Например, ошибка прогноза с +0,45% до -0,28% равна 0,73 п.п., а не 0,73%."),
        ("Nowcast", "Оперативная оценка текущего месяца до выхода официального факта. В данном случае строится по недельным ценам."),
        ("Weekly signal", "Сигнал из недельных цен. Он показывает, что происходит внутри месяца, но не является официальным месячным фактом."),
        ("Последняя неделя как уровень месяца", "Для многих товарных рядов итоговый месячный индекс ближе к сравнению цен последней недели месяца с последней неделей предыдущего месяца, а не к простой сумме недельных изменений."),
        ("SA / сезонно скорректированный ряд", "Ряд, очищенный от типичной календарной сезонности. Если SA-ряд тоже снижается, это признак не только обычного сезона, но и более сильного текущего импульса."),
        ("Компонентный вклад", "Оценка того, сколько компонент добавил к общему ИПЦ: изменение компонента × его вес в корзине."),
        ("Вес компонента", "Доля компонента в потребительской корзине. Чем выше вес, тем сильнее влияние компонента на общий ИПЦ."),
        ("Food shock", "Нестандартное движение продовольственных цен, которое существенно меняет общий прогноз. В мае 2026 это прежде всего огурцы, помидоры и яйца."),
        ("Ансамбль моделей", "Средневзвешенный прогноз нескольких моделей. Он устойчивее одной модели, но может запаздывать при резких изменениях, например при продовольственном шоке или переносе тарифов."),
        ("Экспертная корректировка", "Ручная корректировка модельного прогноза на известные факторы, которые модель могла не учесть: перенос тарифов, свежие недельные цены, изменение курса."),
        ("Сценарный диапазон", "Низкая и высокая границы прогноза при разных вариантах развития событий. Это не ошибка расчёта, а способ показать неопределённость."),
        ("Тарифный перенос", "Изменение календаря индексации регулируемых услуг. В 2026 году июльская индексация не закладывается, основной эффект переносится на октябрь."),
        ("Дезинфляционный фактор", "Фактор, который снижает темп роста цен. Крепкий рубль может временно сдерживать импортные и торгуемые товары."),
        ("Возврат цен / rebound", "Частичный рост после резкого снижения. После падения овощей и яиц возможен обратный ход, поэтому прогноз не должен механически продолжать падение бесконечно."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Термин"
    table.rows[0].cells[1].text = "Простое пояснение"
    for term, explanation in glossary:
        cells = table.add_row().cells
        cells[0].text = term
        cells[1].text = explanation

    doc.save(DOCX_PATH)


def build_zip() -> None:
    files = [
        CSV_PATH,
        DOCX_PATH,
        ROOT / "archive" / "results" / "april_2026_deviation_analysis" / "april_2026_forecast_deviation_analysis.docx",
        ROOT / "archive" / "results" / "april_2026_deviation_analysis" / "april_2026_forecast_deviation_calculations.xlsx",
        OUT / "verification_report.md",
    ]
    with zipfile.ZipFile(ZIP_PATH, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for path in files:
            if path.exists():
                zf.write(path, arcname=path.name)
            else:
                print(f"WARNING missing: {path}")


def main() -> None:
    import argparse
    parser = argparse.ArgumentParser(description="Build the historical or a dated verified forecast package")
    parser.add_argument('--update-dir', type=Path, help='Dated folder with REPORT.md, full_path_2026_2027.csv and verification_report.md')
    parser.add_argument('--historical', action='store_true', help='Explicitly rebuild the original historical hardcoded package')
    args = parser.parse_args()
    pointer = OUT / 'current_update.json'
    if not args.update_dir and not args.historical and pointer.exists():
        import json
        args.update_dir = ROOT / json.loads(pointer.read_text())['update_dir']
    if args.update_dir:
        build_updated_package(args.update_dir.resolve())
        return
    OUT.mkdir(parents=True, exist_ok=True)
    write_csv()
    write_docx()
    build_zip()
    print(CSV_PATH)
    print(DOCX_PATH)
    print(ZIP_PATH)


def build_updated_package(source: Path) -> None:
    """Publish a dated scenario without reusing the historical hardcoded narrative."""
    import re
    import shutil
    import pandas as pd
    from docx.shared import Inches
    required = ['REPORT.md', 'full_path_2026_2027.csv', 'forecast_2026_2027.csv',
                'forecast_calculations.xlsx', 'forecast.png', 'verification_report.md']
    for name in required:
        if not (source / name).is_file():
            raise FileNotFoundError(source / name)
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(10)
    def clean(text):
        return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text).replace('**', '').replace('`', '')
    for block in (source / 'REPORT.md').read_text().split('\n\n'):
        lines = block.strip().splitlines()
        if not lines:
            continue
        if lines[0].startswith('#'):
            level = min(len(lines[0])-len(lines[0].lstrip('#')), 3)
            doc.add_heading(clean(lines[0]).lstrip('# '), level=level)
        elif lines[0].startswith('|'):
            rows = [[clean(c.strip()) for c in line.strip('|').split('|')]
                    for line in lines if not re.fullmatch(r'[| :\-]+', line)]
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            for values in rows:
                for cell, value in zip(table.add_row().cells, values):
                    cell.text = value
        else:
            paragraphs = re.split(r'\n(?=\d+\. )', '\n'.join(lines))
            for paragraph in paragraphs:
                doc.add_paragraph(clean(' '.join(s.strip() for s in paragraph.splitlines())))
    doc.add_heading('Помесячная траектория', level=1)
    frame = pd.read_csv(source / 'full_path_2026_2027.csv')
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    for cell, label in zip(table.rows[0].cells, ['Месяц', 'Статус', 'ИПЦ м/м, индекс', 'Прирост г/г, %']):
        cell.text = label
    for row in frame.itertuples():
        cells = table.add_row().cells
        for cell, value in zip(cells, [str(row.date)[:7], 'Факт' if row.status == 'fact' else 'Прогноз', f'{row.mom_index:.2f}', f'{row.yoy_pct:.2f}']):
            cell.text = value
    doc.add_picture(str(source / 'forecast.png'), width=Inches(6.3))
    docx = source / 'forecast_explanation.docx'
    doc.save(docx)
    files = required + ['forecast_explanation.docx', 'forecast.html', 'calculation_manifest.json',
                        'forecast_explanation.pdf', 'verification.json', 'diagnostics.csv',
                        'august_model_comparison.csv', 'prior_forecast_audit.md', 'tariff_sensitivity.csv',
                        'build_update.py', 'refresh_models.py', 'update_opr.py', 'verify_update.py']
    with zipfile.ZipFile(source / 'forecast_package_20260928.zip', 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        for name in files:
            if (source / name).exists():
                zf.write(source / name, arcname=name)
        zf.write(ROOT / 'assets/06_2026_02_Прогноз.xlsx', arcname='06_2026_02_Прогноз.xlsx')
    # Stable publication paths consumed by the project; preserve the previous package.
    previous = source / 'previous_full_package'
    previous.mkdir(exist_ok=True)
    for src, dest in [(docx, DOCX_PATH), (source / 'forecast_package_20260928.zip', ZIP_PATH),
                      (source / 'verification_report.md', OUT / 'verification_report.md')]:
        if dest.exists() and not (previous / dest.name).exists():
            shutil.copy2(dest, previous / dest.name)
        shutil.copy2(src, dest)
    if CSV_PATH.exists() and not (previous / CSV_PATH.name).exists():
        shutil.copy2(CSV_PATH, previous / CSV_PATH.name)
    export = pd.DataFrame({'Date': pd.to_datetime(frame.date).dt.strftime('%d.%m.%Y'),
                           'MoM': frame.mom_index, 'YoY': frame.yoy_pct + 100})
    export.to_csv(CSV_PATH, sep=';', decimal=',', index=False, float_format='%.2f', encoding='utf-8-sig')
    import json
    (OUT / 'current_update.json').write_text(json.dumps({'update_dir': str(source.relative_to(ROOT))}, indent=2) + '\n')
    print(source / 'forecast_package_20260928.zip')


if __name__ == "__main__":
    main()
