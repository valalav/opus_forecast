#!/usr/bin/env python3
"""Build the portable app's compact regional CPI JSON from the public XLSX release.

This is a maintainer-side import tool. The portable app consumes JSON and does
not require Python. Requires openpyxl only when rebuilding from an XLSX source.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

AGGREGATES = {1: "y", 2: "nonfood", 3: "food", 4: "services"}
SOURCE_URL = "https://disk.yandex.ru/d/i6czrcdNO0I7BQ"
SOURCE_FILE = "2026-08 ИПЦ с исключением сезонности (регионы).xlsx"
SOURCE_ETAG = "1e85996e5e6c78421d6944f1eefab391"
SOURCE_MODIFIED = "2026-09-15T09:58:45+00:00"
RELEASE_PERIOD = "2026-08"


def iso_month(value: Any) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%m/%d/%y %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    raise ValueError(f"Unrecognized source date: {value!r}")


def number(value: Any) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).strip().replace(" ", "").replace(",", "."))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def classify_region(code: int, name: str) -> str:
    if code == 0:
        return "aggregate"
    if name.endswith("ГУ"):
        return "central_bank_district"
    if "федеральный округ" in name.casefold():
        return "federal_district"
    if "(кроме " in name.casefold():
        return "partial_subject"
    return "subject"


def load_region_names(sheet: Any) -> dict[int, str]:
    names: dict[int, str] = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        code, name = row[1], row[2]
        if code is None or name is None:
            continue
        try:
            names[int(code)] = str(name).strip()
        except (TypeError, ValueError):
            continue
    if not names:
        raise ValueError("No region code/name rows found in 'momsa_regions'")
    return names


def load_macros(path: Path) -> dict[str, dict[str, float | None]]:
    macros: dict[str, dict[str, float | None]] = {}
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream, delimiter=";")
        required = {"Date", "Ki", "Ruonia"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"Macro CSV must contain {sorted(required)}")
        for row in reader:
            day = iso_month(row["Date"])[:7] + "-01"
            macros[day] = {"ki": number(row["Ki"]), "ruonia": number(row["Ruonia"])}
    return macros


def build(source_path: Path, macro_path: Path) -> dict[str, Any]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Rebuilding from XLSX requires openpyxl 3.x") from exc

    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    if "db" not in workbook.sheetnames or "momsa_regions" not in workbook.sheetnames:
        raise ValueError("Expected flat 'db' and region lookup 'momsa_regions' sheets")
    names = load_region_names(workbook["momsa_regions"])
    histories: dict[int, dict[str, dict[str, dict[str, float | None]]]] = defaultdict(
        lambda: {"raw": {}, "sa": {}}
    )
    for row in workbook["db"].iter_rows(min_row=2, values_only=True):
        day, item, region, raw, sa = row[:5]
        if item is None or int(item) not in AGGREGATES:
            continue
        region_code = int(region)
        target = AGGREGATES[int(item)]
        date_key = iso_month(day)
        bucket = histories[region_code]
        for kind, value in (("raw", raw), ("sa", sa)):
            current = bucket[kind].setdefault(date_key, {})
            if target in current:
                raise ValueError(f"Duplicate {kind} observation for {region_code}/{date_key}/{target}")
            current[target] = number(value)
    workbook.close()

    if not histories:
        raise ValueError("No aggregate observations (item codes 1–4) found in 'db'")
    unknown = sorted(set(histories) - set(names))
    if unknown:
        raise ValueError(f"No region names for codes: {unknown}")
    macros = load_macros(macro_path)
    all_dates = sorted({day for history in histories.values() for day in history["raw"]})
    release_latest = max(all_dates)
    release_earliest = min(all_dates)
    expected_fields = tuple(AGGREGATES.values())

    def missing_months(observed: list[str]) -> list[str]:
        if not observed:
            return []
        present = set(observed)
        year, month = map(int, observed[0][:7].split("-"))
        end_year, end_month = map(int, observed[-1][:7].split("-"))
        missing = []
        while (year, month) <= (end_year, end_month):
            candidate = f"{year:04d}-{month:02d}-01"
            if candidate not in present:
                missing.append(candidate)
            month += 1
            if month == 13:
                year += 1
                month = 1
        return missing

    def rows_for(history: dict[str, dict[str, float | None]]) -> list[dict[str, Any]]:
        return [
            {
                "date": day,
                **{field: values.get(field) for field in expected_fields},
                "ki": macros.get(day, {}).get("ki"),
                "ruonia": macros.get(day, {}).get("ruonia"),
            }
            for day, values in sorted(history.items())
        ]

    regions = []
    for code in sorted(histories):
        history = histories[code]
        raw_rows = rows_for(history["raw"])
        sa_rows = rows_for(history["sa"])
        raw_dates = [r["date"] for r in raw_rows]
        sa_dates = [r["date"] for r in sa_rows]
        status = (
            "archival_or_incomplete"
            if max(raw_dates) < release_latest
            else "current_limited_history"
            if min(raw_dates) > release_earliest or len(raw_dates) < len(all_dates)
            else "observed_through_release"
        )
        # Per-region status is descriptive: absent months remain visible as gaps.
        regions.append(
            {
                "code": code,
                "name": names[code],
                "kind": classify_region(code, names[code]),
                "status": status,
                "rows": raw_rows,
                "sa_rows": sa_rows,
                "coverage": {
                    "raw": {"from": min(raw_dates), "through": max(raw_dates), "count": len(raw_dates)},
                    "sa": {"from": min(sa_dates), "through": max(sa_dates), "count": len(sa_dates)},
                    "raw_missing_fields": sum(any(r[f] is None for f in expected_fields) for r in raw_rows),
                    "sa_missing_fields": sum(any(r[f] is None for f in expected_fields) for r in sa_rows),
                    "raw_gaps": missing_months(raw_dates),
                    "sa_gaps": missing_months(sa_dates),
                },
            }
        )

    macro_rows = [
        {"date": day, **values}
        for day, values in sorted(macros.items())
        if all_dates[0] <= day <= release_latest
    ]
    return {
        "schema_version": 1,
        "source": {
            "public_folder": SOURCE_URL,
            "file": SOURCE_FILE,
            "sheet": "db",
            "modified": SOURCE_MODIFIED,
            "etag": SOURCE_ETAG,
            "sha256": sha256_file(source_path),
            "release_period": RELEASE_PERIOD,
            "raw_column": "mom",
            "sa_column": "mom_sa",
            "raw_and_sa_same_release": True,
            "unit": "monthly index level (100 means no change)",
        },
        "coverage": {
            "release_latest": release_latest,
            "regions": len(regions),
            "region_month_range": {"from": min(all_dates), "through": release_latest},
            "raw_periods": ["raw"],
            "sa_periods": ["sa"],
        },
        "macro_source": {"file": macro_path.name, "sha256": sha256_file(macro_path), "columns": {"ki": "Ki", "ruonia": "Ruonia"}},
        "macro_rows": macro_rows,
        "regions": regions,
    }


def make_audit(source_path: Path, full_path: Path | None) -> dict[str, Any]:
    audit: dict[str, Any] = {
        "checked_at": "2026-09-28",
        "public_folder": SOURCE_URL,
        "listing_api": "https://cloud-api.yandex.net/v1/disk/public/resources",
        "api_origin_null": {"status": 200, "access_control_allow_origin": "null"},
        "binary_download_origin_null": {"status": 206, "access_control_allow_origin": "*", "range_requests": True},
        "selected_source": {
            "file": SOURCE_FILE,
            "modified": SOURCE_MODIFIED,
            "size_bytes": source_path.stat().st_size,
            "etag": SOURCE_ETAG,
            "sha256": sha256_file(source_path),
            "sheet_contract": {"db": ["Day", "item", "region", "mom", "mom_sa"], "region_lookup": "momsa_regions columns B=code, C=name"},
            "coverage": "item codes 1-4 across 101 region entities; 100 entities have 128 observations (2016-01 through 2026-08), and code 93 Южный федеральный округ has 120 (2016-09 through 2026-08); both mom and mom_sa are present for all exported rows",
        },
        "other_latest_candidate": {
            "file": "2026-08 Статистика ИПЦ полный.xlsx",
            "modified": "2026-09-14T13:39:56+00:00",
            "size_bytes": 105794814,
            "etag": "0a241b5b3e4d0f57eb8b57d954d23758",
            "inspection": "One visible pivot sheet filtered to Россия; pivot cache uses an external Access connection, saveData=0, recordCount=0 and has no pivotCacheRecords XML. Regional data may exist inside embedded Excel Data Model binary, which this importer does not decode.",
        },
        "browser_update": "Yandex public listing and binary download support cross-origin requests; XLSX parsing remains a separate browser implementation. No direct browser ACCDB import is supported or claimed.",
    }
    if full_path and full_path.exists():
        audit["other_latest_candidate"]["sha256"] = sha256_file(full_path)
    return audit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True, help="Downloaded regional SA XLSX release")
    parser.add_argument("--macros", type=Path, required=True, help="Main inflation CSV containing Ki and Ruonia")
    parser.add_argument("--output", type=Path, required=True, help="Output regional_indices.json")
    parser.add_argument("--audit", type=Path, help="Optional source_audit.json output")
    parser.add_argument("--full-workbook", type=Path, help="Optional downloaded full-statistics XLSX for SHA audit")
    args = parser.parse_args()

    package = build(args.source, args.macros)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(package, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    if args.audit:
        args.audit.parent.mkdir(parents=True, exist_ok=True)
        audit = make_audit(args.source, args.full_workbook)
        args.audit.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {args.output}: {len(package['regions'])} regions, {package['coverage']['region_month_range']['from']}..{package['coverage']['release_latest']}")


if __name__ == "__main__":
    main()
